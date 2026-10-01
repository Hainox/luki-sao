"""Свод по люкам: арифметика, формат процента, период по Москве, Excel."""
from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.formatting import card_label, percent_label
from app.services.periods import resolve_period
from app.services.summary import FOOTNOTE, XLSX_HEADER
from tests.conftest import SAO_DISTRICTS, create_card, login_as, run_sql, set_created_at, upload


@pytest.mark.parametrize("fixed,detected,expected", [
    (2, 3, "66,7%"),
    (7, 10, "70%"),
    (0, 0, "—"),
    (0, 4, "0%"),
    (3, 3, "100%"),
    (1, 8, "12,5%"),
    (5, 16, "31,3%"),  # 31,25 → половина вверх, а не банковское округление
    (1, 3, "33,3%"),
])
def test_percent_label(fixed, detected, expected):
    assert percent_label(fixed, detected) == expected


@pytest.mark.parametrize("number,expected", [(1, "ОЛХ-001"), (42, "ОЛХ-042"), (999, "ОЛХ-999"), (1000, "ОЛХ-1000")])
def test_card_label(number, expected):
    assert card_label(number) == expected


def test_rolling_periods_by_moscow_date():
    # 30.09 21:30 UTC — уже 1 октября по Москве.
    now = datetime(2026, 9, 30, 21, 30, tzinfo=timezone.utc)
    assert resolve_period("today", now=now).date_from == date(2026, 10, 1)
    week = resolve_period("week", now=now)
    assert (week.date_from, week.date_to) == (date(2026, 9, 25), date(2026, 10, 1))
    month = resolve_period("month", now=now)
    assert (month.date_from, month.date_to) == (date(2026, 9, 2), date(2026, 10, 1))
    start, end = resolve_period("today", now=now).utc_bounds
    assert start == datetime(2026, 9, 30, 21, 0, tzinfo=timezone.utc)
    assert end == datetime(2026, 10, 1, 21, 0, tzinfo=timezone.utc)
    assert resolve_period("custom", date(2026, 9, 1), None).label == "01.09.2026"
    assert resolve_period("custom", date(2026, 9, 1), date(2026, 9, 7)).label == "01.09.2026 — 07.09.2026"


async def _make_cards(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    sokol = await login_as(client, jj, "inspector2", district="Сокол")
    # Аэропорт: 3 выявлено, 2 принято, 1 на проверке → 66,7%
    for final in ("accepted", "accepted", "on_review"):
        card = await create_card(client, aero)
        await upload(client, aero, card["id"], "after")
        if final == "accepted":
            await client.post(f"/api/cards/{card['id']}/accept", headers=admin)
    # Сокол: 10 выявлено, 7 принято → 70%
    for i in range(10):
        card = await create_card(client, sokol)
        if i < 7:
            await upload(client, sokol, card["id"], "after")
            await client.post(f"/api/cards/{card['id']}/accept", headers=admin)
    return aero


async def test_summary_table(client, jj, admin):
    await _make_cards(client, jj, admin)
    r = await client.get("/api/summary", headers=admin)
    assert r.status_code == 200
    data = r.json()
    assert [row["district_name"] for row in data["rows"]] == sorted(SAO_DISTRICTS)
    by_name = {row["district_name"]: row for row in data["rows"]}
    assert (by_name["Аэропорт"]["detected"], by_name["Аэропорт"]["fixed"],
            by_name["Аэропорт"]["on_review"], by_name["Аэропорт"]["percent_label"]) == (3, 2, 1, "66,7%")
    assert by_name["Сокол"]["percent_label"] == "70%"
    assert by_name["Сокол"]["percent"] == "70.0"
    assert by_name["Коптево"]["detected"] == 0
    assert by_name["Коптево"]["percent_label"] == "—"
    assert by_name["Коптево"]["percent"] is None
    total = data["total"]
    assert total["district_name"] == "Итого по САО"
    assert (total["detected"], total["fixed"], total["on_review"]) == (13, 9, 1)
    assert total["percent_label"] == "69,2%"
    assert data["period"] == {"kind": "all", "date_from": None, "date_to": None, "label": "За всё время"}


async def test_summary_period_uses_moscow_midnight(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    just_after = await create_card(client, aero)
    just_before = await create_card(client, aero)
    set_created_at(just_after["id"], "2026-09-14T21:30:00+00:00")   # 15.09 00:30 МСК
    set_created_at(just_before["id"], "2026-09-14T20:59:00+00:00")  # 14.09 23:59 МСК

    params = {"period": "custom", "date_from": "2026-09-15", "date_to": "2026-09-15"}
    data = (await client.get("/api/summary", params=params, headers=admin)).json()
    assert data["total"]["detected"] == 1
    assert data["period"]["label"] == "15.09.2026"

    params = {"period": "custom", "date_from": "2026-09-14", "date_to": "2026-09-14"}
    assert (await client.get("/api/summary", params=params, headers=admin)).json()["total"]["detected"] == 1


async def test_refix_after_return_does_not_increase_detected(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    card = await create_card(client, aero)
    await upload(client, aero, card["id"], "after")
    await client.post(f"/api/cards/{card['id']}/return", json={"comment": "Переделать"}, headers=admin)
    summary = (await client.get("/api/summary", headers=admin)).json()
    row = next(r for r in summary["rows"] if r["district_name"] == "Аэропорт")
    assert (row["detected"], row["fixed"], row["on_review"], row["percent_label"]) == (1, 0, 0, "0%")

    await upload(client, aero, card["id"], "after")
    summary = (await client.get("/api/summary", headers=admin)).json()
    row = next(r for r in summary["rows"] if r["district_name"] == "Аэропорт")
    assert (row["detected"], row["fixed"], row["on_review"]) == (1, 0, 1)

    await client.post(f"/api/cards/{card['id']}/accept", headers=admin)
    summary = (await client.get("/api/summary", headers=admin)).json()
    row = next(r for r in summary["rows"] if r["district_name"] == "Аэропорт")
    assert (row["detected"], row["fixed"], row["on_review"], row["percent_label"]) == (1, 1, 0, "100%")
    assert summary["total"]["detected"] == 1


@pytest.mark.parametrize("role", ["inspector", "reviewer"])
async def test_okrug_summary_is_prefecture_only(client, jj, admin, role):
    staff = await login_as(client, jj, f"aero_{role}", role=role, district="Аэропорт")
    await create_card(client, staff)
    for path in ("/api/summary", "/api/summary.xlsx"):
        r = await client.get(path, headers=staff)
        assert r.status_code == 403, path
        assert r.json()["detail"] == "Свод по всем районам доступен только префектуре"
        assert (await client.get(path, headers=admin)).status_code == 200


@pytest.mark.parametrize("role", ["inspector", "reviewer"])
async def test_okrug_summary_without_district_explains_why(client, jj, admin, role):
    orphan = await login_as(client, jj, f"orphan_{role}", role=role, district=None)
    for path in ("/api/summary", "/api/summary.xlsx"):
        r = await client.get(path, headers=orphan)
        assert r.status_code == 403, path
        assert "не назначен район" in r.json()["detail"]


async def test_summary_requires_login(client):
    assert (await client.get("/api/summary")).status_code == 401
    assert (await client.get("/api/summary.xlsx")).status_code == 401


async def test_summary_excludes_unknown_district(client, jj, admin):
    names = [r["district_name"] for r in (await client.get("/api/summary", headers=admin)).json()["rows"]]
    assert "Неизвестный район" not in names
    assert len(names) == 16


async def test_summary_xlsx(client, jj, admin):
    await _make_cards(client, jj, admin)
    r = await client.get("/api/summary.xlsx", headers=admin)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    ws = load_workbook(BytesIO(r.content)).active
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    assert rows[0][0] == "Свод по люкам САО — неудовлетворительные ОЛХ"
    assert rows[1][0] == "Период: За всё время"
    assert rows[2] == XLSX_HEADER
    body = {row[0]: row for row in rows[3:3 + 17]}
    assert body["Аэропорт"] == ["Аэропорт", 3, 2, 1, "66,7%"]
    assert body["Коптево"] == ["Коптево", 0, 0, 0, "—"]
    assert body["Итого по САО"] == ["Итого по САО", 13, 9, 1, "69,2%"]
    assert rows[-1][0] == FOOTNOTE


async def test_summary_xlsx_is_formula_injection_safe(client, jj, admin):
    run_sql("UPDATE districts SET name = %s WHERE name = 'Коптево'", ('=HYPERLINK("http://evil.test","x")',))
    r = await client.get("/api/summary.xlsx", params={"period": "month"}, headers=admin)
    ws = load_workbook(BytesIO(r.content)).active
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(cells) == 1
    assert cells[0].data_type == "s"

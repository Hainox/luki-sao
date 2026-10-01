"""Справочник ДТ/ОДХ: выбор места люка, подсказка «рядом с вами», фильтр
«Все / ДТ / ОДХ» в журнале и сводах."""

import gzip
import json
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

from app.load_territories import DEFAULT_PATH, load_dataset
from app.services.territories import distance_m, district_key
from tests.conftest import (
    DISTRICT_IDS,
    SAO_DISTRICTS,
    TERRITORY_IDS,
    create_card,
    district_corner,
    login_as,
    place,
    run_sql,
    territories_dataset,
)


async def test_district_staff_get_own_dt_and_odh(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    r = await client.get("/api/territories", headers=aero)
    assert r.status_code == 200
    assert [(t["kind"], t["name"]) for t in r.json()] == [
        ("dt", "Двор района Аэропорт"),
        ("odh", "Улица района Аэропорт"),
    ]
    assert r.json()[0]["owner"] == "Жилищник Аэропорт"

    r = await client.get(
        "/api/territories", params={"district_id": DISTRICT_IDS["Сокол"]}, headers=aero
    )
    assert r.status_code == 403

    r = await client.get("/api/territories", headers=admin)
    assert r.status_code == 422
    r = await client.get(
        "/api/territories", params={"district_id": DISTRICT_IDS["Сокол"]}, headers=admin
    )
    assert [t["name"] for t in r.json()] == ["Двор района Сокол", "Улица района Сокол"]

    orphan = await login_as(client, jj, "orphan", district=None)
    assert (await client.get("/api/territories", headers=orphan)).status_code == 403


async def test_registry_names_without_yo_match_districts(client, jj):
    # В журнале обходов «Савёловский», в реестре — «Савеловский».
    savel = await login_as(client, jj, "inspector1", district="Савёловский")
    r = await client.get("/api/territories", headers=savel)
    assert [t["name"] for t in r.json()] == ["Двор района Савёловский", "Улица района Савёловский"]


async def test_nearby_sorts_by_distance_within_own_district(client, jj):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    lat, lon = district_corner("Аэропорт")

    inside_dt = {"lat": lat + 0.0005, "lon": lon + 0.0005}
    r = await client.get("/api/territories/nearby", params=inside_dt, headers=aero)
    assert r.status_code == 200
    found = r.json()
    assert [(t["kind"], t["distance_m"]) for t in found][0] == ("dt", 0)
    # ОДХ — в 0,001° (~63 м) к востоку от края ДТ.
    assert found[1]["kind"] == "odh"
    assert 85 <= found[1]["distance_m"] <= 100

    # Соседний район (Беговой) — в 0,01° севернее: его объекты не предлагаем.
    between = {"lat": lat + 0.0055, "lon": lon + 0.0005}
    r = await client.get("/api/territories/nearby", params=between, headers=aero)
    assert all(t["name"].endswith("Аэропорт") for t in r.json())

    far = {"lat": lat - 0.01, "lon": lon}
    assert (await client.get("/api/territories/nearby", params=far, headers=aero)).json() == []


async def test_card_stores_place_and_rejects_mismatches(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    card = await create_card(client, aero, address="у подъезда 2", kind="odh")
    assert card["place_kind"] == "odh"
    assert card["territory"]["id"] == TERRITORY_IDS[("Аэропорт", "odh")]
    assert card["territory"]["passport_url"] == "https://reestr-ogh.mos.ru/ogh/odh-Аэропорт"
    assert card["address"] == "Улица района Аэропорт — у подъезда 2"
    plain = await create_card(client, aero, address=None)
    assert plain["address"] == "Двор района Аэропорт"

    r = await client.post("/api/cards", json={"address_note": "ул. Зорге, 1"}, headers=aero)
    assert r.status_code == 422

    wrong_kind = {"place_kind": "dt", "territory_id": TERRITORY_IDS[("Аэропорт", "odh")]}
    r = await client.post("/api/cards", json=wrong_kind, headers=aero)
    assert r.status_code == 422
    assert r.json()["detail"] == "«Улица района Аэропорт» — это ОДХ, а выбрано ДТ"

    r = await client.post("/api/cards", json=place("Сокол"), headers=aero)
    assert r.status_code == 422
    assert r.json()["detail"] == "«Двор района Сокол» относится к другому району"

    r = await client.post(
        "/api/cards",
        json={**place("Сокол"), "district_id": DISTRICT_IDS["Аэропорт"]},
        headers=admin,
    )
    assert r.status_code == 422

    run_sql(
        "UPDATE territories SET is_active = FALSE WHERE id = %s",
        (TERRITORY_IDS[("Аэропорт", "dt")],),
    )
    try:
        r = await client.post("/api/cards", json=place("Аэропорт"), headers=aero)
        assert r.status_code == 422
        assert "не найден в справочнике" in r.json()["detail"]
    finally:
        run_sql(
            "UPDATE territories SET is_active = TRUE WHERE id = %s",
            (TERRITORY_IDS[("Аэропорт", "dt")],),
        )


async def test_place_filter_in_journal_and_summaries(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    await create_card(client, aero, kind="dt")
    await create_card(client, aero, kind="dt")
    odh = await create_card(client, aero, kind="odh")
    legacy = await create_card(client, aero)
    run_sql(
        "UPDATE cards SET place_kind = NULL, territory_id = NULL WHERE id = %s", (legacy["id"],)
    )

    r = (await client.get("/api/cards", params={"place": "odh"}, headers=aero)).json()
    assert [c["id"] for c in r["items"]] == [odh["id"]]
    assert r["counts"]["all"] == 1
    assert (await client.get("/api/cards", params={"place": "dt"}, headers=aero)).json()["counts"][
        "all"
    ] == 2
    assert (await client.get("/api/cards", headers=aero)).json()["counts"]["all"] == 4

    def aero_row(summary):
        return next(row for row in summary["rows"] if row["district_name"] == "Аэропорт")

    assert aero_row((await client.get("/api/summary", headers=admin)).json())["detected"] == 4
    summary = (await client.get("/api/summary", params={"place": "dt"}, headers=admin)).json()
    assert summary["place"] == "dt"
    assert aero_row(summary)["detected"] == 2
    assert summary["total"]["detected"] == 2

    district = (
        await client.get("/api/summary/district", params={"place": "odh"}, headers=aero)
    ).json()
    assert district["place"] == "odh"
    assert district["totals"]["detected"] == 1
    assert [c["id"] for c in district["oldest_open"]] == [odh["id"]]

    r = await client.get("/api/summary.xlsx", params={"place": "odh"}, headers=admin)
    assert quote("Свод по люкам САО (ОДХ).xlsx") in r.headers["content-disposition"]
    rows = [row[0] for row in load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True)]
    assert "Где найдены: только объекты дорожного хозяйства (ОДХ)" in rows

    r = await client.get("/api/summary/district.xlsx", params={"place": "dt"}, headers=aero)
    assert quote("Свод по люкам — Аэропорт (ДТ).xlsx") in r.headers["content-disposition"]
    rows = [row[0] for row in load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True)]
    assert "Где найдены: только дворовые территории (ДТ)" in rows

    assert (
        await client.get("/api/cards", params={"place": "lyuk"}, headers=aero)
    ).status_code == 422


async def test_reload_updates_deactivates_and_skips_unchanged(client):
    from sqlalchemy.ext.asyncio import AsyncSession

    from app.database import get_engine

    def async_session():
        return AsyncSession(get_engine(), expire_on_commit=False)

    original = gzip.compress(json.dumps(territories_dataset(), ensure_ascii=False).encode())
    changed = territories_dataset()
    changed["items"] = [i for i in changed["items"] if i["registry_id"] != "odh-Сокол"]
    changed["items"][0]["name"] = "Двор района Аэропорт (новая версия паспорта)"
    try:
        async with async_session() as db:
            message = await load_dataset(
                db, gzip.compress(json.dumps(changed, ensure_ascii=False).encode())
            )
        assert "ДТ 16, ОДХ 15" in message
        rows = dict(run_sql("SELECT registry_id, is_active FROM territories"))
        assert rows["odh-Сокол"] is False
        assert rows["odh-Аэропорт"] is True
        name = run_sql(
            "SELECT name FROM territories WHERE id = %s", (TERRITORY_IDS[("Аэропорт", "dt")],)
        )
        assert name == [("Двор района Аэропорт (новая версия паспорта)",)]
    finally:
        async with async_session() as db:
            await load_dataset(db, original)
    async with async_session() as db:
        assert await load_dataset(db, original) == "Справочник ДТ/ОДХ не менялся"
    assert dict(run_sql("SELECT registry_id, is_active FROM territories"))["odh-Сокол"] is True


def test_distance_respects_holes():
    outer = [[37.0, 55.0], [37.01, 55.0], [37.01, 55.01], [37.0, 55.01], [37.0, 55.0]]
    hole = [
        [37.004, 55.004],
        [37.006, 55.004],
        [37.006, 55.006],
        [37.004, 55.006],
        [37.004, 55.004],
    ]
    assert distance_m([[outer, hole]], 55.002, 37.002) == 0
    # Центр дыры — во дворе внутри ОДХ-кольца, до края дыры ~63 м по долготе.
    assert 60 < distance_m([[outer, hole]], 55.005, 37.005) < 70


def test_shipped_dataset_covers_every_district():
    dataset = json.loads(gzip.decompress(Path(DEFAULT_PATH).read_bytes()))
    items = dataset["items"]
    assert sum(i["kind"] == "dt" for i in items) == 2131
    assert sum(i["kind"] == "odh" for i in items) == 689
    for name in SAO_DISTRICTS:
        kinds = {
            i["kind"]
            for i in items
            if district_key(name) in {district_key(d) for d in i["districts"]}
        }
        assert kinds == {"dt", "odh"}, name
    assert all(len(ring) >= 4 for i in items for polygon in i["polygons"] for ring in polygon)
    assert not any(i["name"].startswith("ДТ\\") for i in items)

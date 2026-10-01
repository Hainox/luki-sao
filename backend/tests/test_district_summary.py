"""Подробный свод по району: права, показатели, динамика, давние нарушения, Excel."""

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import quote

import pytest
from openpyxl import load_workbook

from app.services.district_summary import dynamics_buckets
from app.services.periods import MSK, msk_today
from app.services.summary import FOOTNOTE
from tests.conftest import DISTRICT_IDS, create_card, login_as, run_sql, set_created_at, upload

DISTRICT_PATHS = ("/api/summary/district", "/api/summary/district.xlsx")


def test_dynamics_by_day_up_to_31_days():
    unit, buckets = dynamics_buckets(date(2026, 9, 1), date(2026, 10, 1))
    assert unit == "day"
    assert len(buckets) == 31
    assert buckets[0] == (date(2026, 9, 1), date(2026, 9, 1), "01.09")
    assert buckets[-1] == (date(2026, 10, 1), date(2026, 10, 1), "01.10")


def test_dynamics_by_month_when_longer():
    unit, buckets = dynamics_buckets(date(2026, 1, 15), date(2026, 2, 15))  # 32 дня
    assert unit == "month"
    assert buckets == [
        (date(2026, 1, 15), date(2026, 1, 31), "январь 2026"),
        (date(2026, 2, 1), date(2026, 2, 15), "февраль 2026"),
    ]
    unit, buckets = dynamics_buckets(date(2025, 11, 3), date(2026, 2, 28))
    assert [b[2] for b in buckets] == ["ноябрь 2025", "декабрь 2025", "январь 2026", "февраль 2026"]
    assert buckets[-1][:2] == (date(2026, 2, 1), date(2026, 2, 28))


# ── Права ───────────────────────────────────────────────────────


@pytest.mark.parametrize("role", ["inspector", "reviewer"])
async def test_district_staff_get_own_district_only(client, jj, admin, role):
    staff = await login_as(client, jj, f"aero_{role}", role=role, district="Аэропорт")
    for path in DISTRICT_PATHS:
        r = await client.get(path, headers=staff)
        assert r.status_code == 200, (path, r.text)
        r = await client.get(path, params={"district_id": DISTRICT_IDS["Аэропорт"]}, headers=staff)
        assert r.status_code == 200, (path, r.text)
        r = await client.get(path, params={"district_id": DISTRICT_IDS["Сокол"]}, headers=staff)
        assert r.status_code == 403, path
        assert r.json()["detail"] == "Статистика другого района вам недоступна"
    data = (await client.get("/api/summary/district", headers=staff)).json()
    assert data["district"] == {"id": DISTRICT_IDS["Аэропорт"], "name": "Аэропорт"}


@pytest.mark.parametrize("role", ["inspector", "reviewer"])
async def test_staff_without_district_get_no_district_summary(client, jj, admin, role):
    orphan = await login_as(client, jj, f"orphan_{role}", role=role, district=None)
    for path in DISTRICT_PATHS:
        for params in ({}, {"district_id": DISTRICT_IDS["Аэропорт"]}):
            r = await client.get(path, params=params, headers=orphan)
            assert r.status_code == 403, (path, params)
            assert "не назначен район" in r.json()["detail"]
            assert "администратору журнала обходов" in r.json()["detail"]


async def test_prefecture_opens_any_district_but_must_choose_one(client, jj, admin):
    for path in DISTRICT_PATHS:
        r = await client.get(path, headers=admin)
        assert r.status_code == 422, path
        assert r.json()["detail"] == "Выберите район"
        for name in ("Аэропорт", "Сокол"):
            r = await client.get(path, params={"district_id": DISTRICT_IDS[name]}, headers=admin)
            assert r.status_code == 200, (path, name)
    data = (
        await client.get(
            "/api/summary/district", params={"district_id": DISTRICT_IDS["Сокол"]}, headers=admin
        )
    ).json()
    assert data["district"]["name"] == "Сокол"


async def test_district_summary_unknown_district_and_login(client, jj, admin):
    r = await client.get(
        "/api/summary/district",
        params={"district_id": "00000000-0000-0000-0000-000000000000"},
        headers=admin,
    )
    assert r.status_code == 404
    for path in DISTRICT_PATHS:
        assert (await client.get(path)).status_code == 401


# ── Показатели ──────────────────────────────────────────────────


async def _accept(client, staff, admin, card_id):
    await upload(client, staff, card_id, "after")
    r = await client.post(f"/api/cards/{card_id}/accept", headers=admin)
    assert r.status_code == 200


async def _return(client, staff, admin, card_id):
    await upload(client, staff, card_id, "after")
    r = await client.post(
        f"/api/cards/{card_id}/return", json={"comment": "Переделать"}, headers=admin
    )
    assert r.status_code == 200


async def test_totals(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    sokol = await login_as(client, jj, "inspector2", district="Сокол")

    fast = await create_card(client, aero)
    await _accept(client, aero, admin, fast["id"])
    slow = await create_card(client, aero)
    await _return(client, aero, admin, slow["id"])
    await _accept(client, aero, admin, slow["id"])
    review = await create_card(client, aero)
    await upload(client, aero, review["id"], "after")
    twice = await create_card(client, aero)
    await _return(client, aero, admin, twice["id"])
    await _return(client, aero, admin, twice["id"])
    await create_card(client, aero)  # просто выявлено
    other = await create_card(client, sokol)
    await _return(client, sokol, admin, other["id"])

    run_sql(
        "UPDATE cards SET created_at = accepted_at - interval '2 days' WHERE id = %s", (fast["id"],)
    )
    run_sql(
        "UPDATE cards SET created_at = accepted_at - interval '3 days' WHERE id = %s", (slow["id"],)
    )

    data = (await client.get("/api/summary/district", headers=aero)).json()
    assert data["totals"] == {
        "detected": 5,
        "accepted": 2,
        "on_review": 1,
        "open": 2,
        "returned_now": 1,
        "percent_text": "40%",
        "returns_count": 3,
        "avg_days_to_accept": 2.5,
    }
    assert data["period"] == {
        "kind": "all",
        "date_from": None,
        "date_to": None,
        "label": "За всё время",
    }


async def test_totals_follow_period(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    old = await create_card(client, aero)
    await _return(client, aero, admin, old["id"])
    await _accept(client, aero, admin, old["id"])
    set_created_at(old["id"], "2026-08-01T09:00:00+00:00")
    fresh = await create_card(client, aero)
    set_created_at(fresh["id"], "2026-09-10T09:00:00+00:00")

    params = {"period": "custom", "date_from": "2026-09-01", "date_to": "2026-09-30"}
    totals = (await client.get("/api/summary/district", params=params, headers=aero)).json()[
        "totals"
    ]
    assert (totals["detected"], totals["accepted"], totals["open"]) == (1, 0, 1)
    assert totals["returns_count"] == 0
    assert totals["avg_days_to_accept"] is None
    assert totals["percent_text"] == "0%"


async def test_nothing_detected_shows_dash(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    data = (await client.get("/api/summary/district", headers=aero)).json()
    assert data["totals"] == {
        "detected": 0,
        "accepted": 0,
        "on_review": 0,
        "open": 0,
        "returned_now": 0,
        "percent_text": "—",
        "returns_count": 0,
        "avg_days_to_accept": None,
    }
    # За всё время без карточек динамики нет, а у недели — семь пустых дней.
    assert data["dynamics"] == []
    assert data["oldest_open"] == []
    week = (
        await client.get("/api/summary/district", params={"period": "week"}, headers=aero)
    ).json()
    assert week["dynamics_unit"] == "day"
    assert len(week["dynamics"]) == 7
    assert {(b["detected"], b["percent_text"]) for b in week["dynamics"]} == {(0, "—")}


# ── Динамика ────────────────────────────────────────────────────


async def test_dynamics_by_moscow_day(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    after_midnight = await create_card(client, aero)
    before_midnight = await create_card(client, aero)
    await _accept(client, aero, admin, after_midnight["id"])
    set_created_at(after_midnight["id"], "2026-09-02T21:30:00+00:00")  # 03.09 00:30 МСК
    set_created_at(before_midnight["id"], "2026-09-02T20:59:00+00:00")  # 02.09 23:59 МСК

    params = {"period": "custom", "date_from": "2026-09-01", "date_to": "2026-09-07"}
    data = (await client.get("/api/summary/district", params=params, headers=aero)).json()
    assert data["dynamics_unit"] == "day"
    assert [b["label"] for b in data["dynamics"]] == [f"0{d}.09" for d in range(1, 8)]
    by_label = {b["label"]: b for b in data["dynamics"]}
    assert by_label["02.09"] == {
        "label": "02.09",
        "date_from": "2026-09-02",
        "date_to": "2026-09-02",
        "detected": 1,
        "accepted": 0,
        "percent_text": "0%",
    }
    assert (
        by_label["03.09"]["detected"],
        by_label["03.09"]["accepted"],
        by_label["03.09"]["percent_text"],
    ) == (1, 1, "100%")
    assert by_label["01.09"]["percent_text"] == "—"


async def test_dynamics_by_month(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    sokol = await login_as(client, jj, "inspector2", district="Сокол")
    when = {
        "2026-01-20T09:00:00+00:00": "accepted",
        "2026-02-03T09:00:00+00:00": "accepted",
        "2026-02-25T09:00:00+00:00": None,
        "2026-03-10T20:00:00+00:00": None,  # 10.03 23:00 МСК — ещё в периоде
        "2026-03-10T21:30:00+00:00": None,  # 11.03 00:30 МСК — уже нет
    }
    for iso, final in when.items():
        card = await create_card(client, aero)
        if final == "accepted":
            await _accept(client, aero, admin, card["id"])
        set_created_at(card["id"], iso)
    foreign = await create_card(client, sokol)
    set_created_at(foreign["id"], "2026-02-10T09:00:00+00:00")

    params = {"period": "custom", "date_from": "2026-01-15", "date_to": "2026-03-10"}
    data = (await client.get("/api/summary/district", params=params, headers=aero)).json()
    assert data["dynamics_unit"] == "month"
    assert data["dynamics"] == [
        {
            "label": "январь 2026",
            "date_from": "2026-01-15",
            "date_to": "2026-01-31",
            "detected": 1,
            "accepted": 1,
            "percent_text": "100%",
        },
        {
            "label": "февраль 2026",
            "date_from": "2026-02-01",
            "date_to": "2026-02-28",
            "detected": 2,
            "accepted": 1,
            "percent_text": "50%",
        },
        {
            "label": "март 2026",
            "date_from": "2026-03-01",
            "date_to": "2026-03-10",
            "detected": 1,
            "accepted": 0,
            "percent_text": "0%",
        },
    ]
    assert data["totals"]["detected"] == 4


async def test_dynamics_all_time_starts_at_first_card(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    today = msk_today()
    card = await create_card(client, aero)
    first_day = today - timedelta(days=5)
    set_created_at(
        card["id"],
        datetime(first_day.year, first_day.month, first_day.day, 12, tzinfo=MSK).isoformat(),
    )
    data = (await client.get("/api/summary/district", headers=aero)).json()
    assert data["dynamics_unit"] == "day"
    assert [b["date_from"] for b in data["dynamics"]] == [
        (first_day + timedelta(days=i)).isoformat() for i in range(6)
    ]
    assert data["dynamics"][0]["detected"] == 1

    long_ago = today - timedelta(days=100)
    set_created_at(
        card["id"],
        datetime(long_ago.year, long_ago.month, long_ago.day, 12, tzinfo=MSK).isoformat(),
    )
    data = (await client.get("/api/summary/district", headers=aero)).json()
    assert data["dynamics_unit"] == "month"
    assert data["dynamics"][0]["date_from"] == long_ago.isoformat()
    assert data["dynamics"][0]["detected"] == 1
    assert data["dynamics"][-1]["date_to"] == today.isoformat()
    assert sum(b["detected"] for b in data["dynamics"]) == 1


# ── Дольше всех ждут исправления ────────────────────────────────


async def test_oldest_open_ignores_period_and_keeps_ten_oldest(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    sokol = await login_as(client, jj, "inspector2", district="Сокол")
    base = datetime(2026, 6, 1, 9, 0, tzinfo=timezone.utc)

    # Самые старые — уже исправлены или на проверке: в список не попадают.
    accepted = await create_card(client, aero)
    await _accept(client, aero, admin, accepted["id"])
    set_created_at(accepted["id"], (base - timedelta(days=30)).isoformat())
    on_review = await create_card(client, aero)
    await upload(client, aero, on_review["id"], "after")
    set_created_at(on_review["id"], (base - timedelta(days=20)).isoformat())
    foreign = await create_card(client, sokol)
    set_created_at(foreign["id"], (base - timedelta(days=40)).isoformat())

    open_cards = []
    for i in range(12):
        card = await create_card(client, aero, address=f"ул. Усиевича, д. {i + 1}")
        set_created_at(card["id"], (base + timedelta(days=i)).isoformat())
        open_cards.append(card)
    await _return(client, aero, admin, open_cards[1]["id"])

    data = (
        await client.get("/api/summary/district", params={"period": "today"}, headers=aero)
    ).json()
    assert data["totals"]["detected"] == 0
    oldest = data["oldest_open"]
    assert [c["id"] for c in oldest] == [c["id"] for c in open_cards[:10]]
    assert oldest[0]["label"] == open_cards[0]["label"]
    assert oldest[0]["address"] == "Двор района Аэропорт — ул. Усиевича, д. 1"
    assert oldest[0]["place_kind"] == "dt"
    assert oldest[0]["status"] == "detected"
    assert oldest[1]["status"] == "returned"
    today = msk_today()
    assert oldest[0]["age_days"] == (today - base.astimezone(MSK).date()).days
    assert oldest[0]["age_days"] - oldest[9]["age_days"] == 9


# ── Excel ───────────────────────────────────────────────────────


async def test_district_xlsx(client, jj, admin):
    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    done = await create_card(client, aero)
    await _accept(client, aero, admin, done["id"])
    waiting = await create_card(client, aero, address="ул. Зорге, д. 1")

    r = await client.get("/api/summary/district.xlsx", headers=aero)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert r.headers["content-disposition"] == (
        'attachment; filename="luki-sao-svod-rayona.xlsx"; '
        f"filename*=UTF-8''{quote('Свод по люкам — Аэропорт.xlsx')}"
    )
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    assert ws.title == "Свод по району"
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    first = [row[0] for row in rows]
    assert rows[0][0] == "Свод по люкам — Аэропорт"
    assert rows[1][0] == "Период: За всё время"

    kpi = {
        row[0]: row[1]
        for row in rows[first.index("Показатель") + 1 : first.index("Показатель") + 8]
    }
    assert kpi["Выявлено (неудовлетворительные ОЛХ)"] == 2
    assert kpi["Исправлено"] == 1
    assert kpi["Не исправлено (выявлено и возвращено)"] == 1
    assert kpi["% исправления"] == "50%"
    assert kpi["Возвратов на доработку"] == 0
    assert kpi["Среднее время до приёмки, дней"] == 0

    dyn = first.index("День")
    assert rows[dyn][:4] == ["День", "Выявлено", "Исправлено", "% исправления"]
    assert rows[dyn + 1][0] == msk_today().strftime("%d.%m.%Y")
    assert rows[dyn + 1][1:4] == [2, 1, "50%"]

    oldest = first.index("Карточка")
    assert rows[oldest] == ["Карточка", "Адрес", "Статус", "Выявлено", "Ждёт, дней"]
    assert rows[oldest + 1] == [
        waiting["label"],
        "ДТ · Двор района Аэропорт — ул. Зорге, д. 1",
        "Выявлено",
        msk_today().strftime("%d.%m.%Y"),
        0,
    ]
    assert FOOTNOTE in first

    r = await client.get(
        "/api/summary/district.xlsx",
        params={
            "district_id": DISTRICT_IDS["Аэропорт"],
            "period": "custom",
            "date_from": "2026-09-01",
            "date_to": "2026-09-07",
        },
        headers=admin,
    )
    assert r.status_code == 200
    assert (
        quote("Свод по люкам — Аэропорт 01.09.2026 — 07.09.2026.xlsx")
        in r.headers["content-disposition"]
    )


async def test_district_xlsx_empty_and_formula_injection_safe(client, jj, admin):
    sokol = await login_as(client, jj, "inspector2", district="Сокол")
    r = await client.get("/api/summary/district.xlsx", headers=sokol)
    first = [row[0] for row in load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True)]
    assert "Карточек пока нет" in first
    assert "Неисправленных карточек нет" in first

    aero = await login_as(client, jj, "inspector1", district="Аэропорт")
    # Новые карточки начинаются с названия объекта справочника, а у карточек,
    # заведённых до него, адрес — свободный текст сотрудника.
    legacy = await create_card(client, aero)
    run_sql(
        "UPDATE cards SET address = %s, place_kind = NULL, territory_id = NULL WHERE id = %s",
        ('=HYPERLINK("http://evil.test","x")', legacy["id"]),
    )
    r = await client.get("/api/summary/district.xlsx", headers=aero)
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    risky = [
        c
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert [c.value for c in risky] == ['=HYPERLINK("http://evil.test","x")']
    assert all(c.data_type == "s" for c in risky)

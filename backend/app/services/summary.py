"""Свод по люкам: выявлено / исправлено / на проверке по районам за период."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Card
from app.schemas import PeriodOut, SummaryOut, SummaryRow
from app.services.districts import okrug_districts
from app.services.formatting import fixed_percent, percent_label
from app.services.periods import Period
from app.services.xlsx_style import (
    safe_append,
    style_data_row,
    style_header_row,
    style_merged_label,
)

TOTAL_LABEL = "Итого по САО"
XLSX_TITLE = "Свод по люкам САО — неудовлетворительные ОЛХ"
XLSX_HEADER = [
    "Район",
    "Выявлено (неудовлетворительные ОЛХ)",
    "Исправлено",
    "На проверке",
    "% исправления",
]
FOOTNOTE = (
    "Исправлено — принято администратором префектуры. "
    "Процент = исправлено / выявлено × 100. При отсутствии нарушений — прочерк."
)


def _row(district_id, name: str, detected: int, fixed: int, on_review: int) -> SummaryRow:
    return SummaryRow(
        district_id=district_id,
        district_name=name,
        detected=detected,
        fixed=fixed,
        on_review=on_review,
        percent=fixed_percent(fixed, detected),
        percent_label=percent_label(fixed, detected),
    )


async def build_summary(db: AsyncSession, period: Period) -> SummaryOut:
    """«Выявлено» — карточки, созданные в периоде (повторное исправление
    после возврата новой карточки не создаёт и выявленное не увеличивает);
    «Исправлено» и «На проверке» — те из них, что сейчас в этом статусе."""
    conds = []
    start, end = period.utc_bounds
    if start is not None:
        conds.append(Card.created_at >= start)
    if end is not None:
        conds.append(Card.created_at < end)

    stats = {
        r.district_id: r
        for r in (
            await db.execute(
                select(
                    Card.district_id,
                    func.count().label("detected"),
                    func.count().filter(Card.status == "accepted").label("fixed"),
                    func.count().filter(Card.status == "on_review").label("on_review"),
                )
                .where(*conds)
                .group_by(Card.district_id)
            )
        ).all()
    }

    rows = []
    for d in await okrug_districts(db):
        s = stats.get(d.id)
        rows.append(
            _row(
                d.id,
                d.name,
                s.detected if s else 0,
                s.fixed if s else 0,
                s.on_review if s else 0,
            )
        )

    total = _row(
        None,
        TOTAL_LABEL,
        sum(r.detected for r in rows),
        sum(r.fixed for r in rows),
        sum(r.on_review for r in rows),
    )
    return SummaryOut(
        period=PeriodOut(
            kind=period.kind,
            date_from=period.date_from,
            date_to=period.date_to,
            label=period.label,
        ),
        rows=rows,
        total=total,
    )


def summary_xlsx(summary: SummaryOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Свод по люкам"
    ncols = len(XLSX_HEADER)

    safe_append(ws, [XLSX_TITLE])
    style_merged_label(ws, ws.max_row, ncols, header=True, fill=False)
    safe_append(ws, [f"Период: {summary.period.label}"])
    style_merged_label(ws, ws.max_row, ncols)

    safe_append(ws, XLSX_HEADER)
    style_header_row(ws, ws.max_row, ncols)

    for r in [*summary.rows, summary.total]:
        safe_append(ws, [r.district_name, r.detected, r.fixed, r.on_review, r.percent_label])
        style_data_row(ws, ws.max_row, ncols)
    for col in range(1, ncols + 1):
        ws.cell(ws.max_row, col).font = Font(bold=True)

    ws.append([])
    safe_append(ws, [FOOTNOTE])
    style_merged_label(ws, ws.max_row, ncols)
    ws.row_dimensions[ws.max_row].height = 32

    for letter, width in zip("ABCDE", (30, 22, 14, 14, 16)):
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[3].height = 45

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

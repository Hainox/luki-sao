"""Подробный свод по одному району: показатели, динамика, давние нарушения."""

import calendar
from datetime import date, datetime, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Literal, cast

from openpyxl import Workbook
from sqlalchemy import Date, extract, func, literal_column, select
from sqlalchemy import cast as sa_cast
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Card, CardEvent, District
from app.schemas import (
    DistrictOut,
    DistrictSummaryOut,
    DistrictSummaryTotals,
    DynamicsBucket,
    OldestOpenCard,
    PeriodOut,
)
from app.services.formatting import card_label, percent_label
from app.services.periods import MSK, Period, msk_today
from app.services.summary import FOOTNOTE
from app.services.xlsx_style import (
    safe_append,
    style_data_row,
    style_header_row,
    style_merged_label,
)

OPEN_STATUSES = ("detected", "returned")
OLDEST_OPEN_LIMIT = 10
# До месяца включительно динамика по дням, длиннее — по месяцам: 31 строка
# ещё читается на телефоне, а «За всё время» по дням — уже нет.
MAX_DAY_BUCKETS = 31

MONTHS = (
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)
STATUS_TEXT = {
    "detected": "Выявлено",
    "on_review": "На проверке",
    "accepted": "Принято",
    "returned": "Возвращено",
}
PERIOD_NOTE = (
    "Показатели и динамика — по карточкам, выявленным в периоде; «Дольше всех ждут "
    "исправления» — все неисправленные карточки района, независимо от периода."
)

# Дата создания по Москве — та же граница суток, что у фильтра периода.
# Пояс — литералом, а не параметром: с параметром asyncpg отправит в SELECT
# и в GROUP BY разные $1/$2, и Postgres не признает их одним выражением.
_MSK_DAY = sa_cast(func.timezone(literal_column("'Europe/Moscow'"), Card.created_at), Date)


def _period_conds(district_id, period: Period) -> list:
    conds = [Card.district_id == district_id]
    start, end = period.utc_bounds
    if start is not None:
        conds.append(Card.created_at >= start)
    if end is not None:
        conds.append(Card.created_at < end)
    return conds


def _days(seconds) -> float | None:
    if seconds is None:
        return None
    days = Decimal(seconds) / Decimal(86400)
    return float(days.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))


DynamicsUnit = Literal["day", "month"]


def dynamics_buckets(start: date, end: date) -> tuple[DynamicsUnit, list[tuple[date, date, str]]]:
    """Интервалы динамики [с, по] включительно, обрезанные по границам периода."""
    if (end - start).days + 1 <= MAX_DAY_BUCKETS:
        days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
        return "day", [(d, d, d.strftime("%d.%m")) for d in days]
    buckets = []
    cur = start
    while cur <= end:
        month_end = cur.replace(day=calendar.monthrange(cur.year, cur.month)[1])
        buckets.append((cur, min(month_end, end), f"{MONTHS[cur.month - 1]} {cur.year}"))
        cur = month_end + timedelta(days=1)
    return "month", buckets


async def _totals(db: AsyncSession, conds: list) -> DistrictSummaryTotals:
    row = (
        await db.execute(
            select(
                func.count().label("detected"),
                func.count().filter(Card.status == "accepted").label("accepted"),
                func.count().filter(Card.status == "on_review").label("on_review"),
                func.count().filter(Card.status.in_(OPEN_STATUSES)).label("open"),
                func.count().filter(Card.status == "returned").label("returned_now"),
                func.avg(extract("epoch", Card.accepted_at - Card.created_at))
                .filter(Card.status == "accepted", Card.accepted_at.is_not(None))
                .label("accept_seconds"),
            ).where(*conds)
        )
    ).one()
    returns_count = (
        await db.execute(
            select(func.count())
            .select_from(CardEvent)
            .join(Card, Card.id == CardEvent.card_id)
            .where(CardEvent.kind == "returned", *conds)
        )
    ).scalar_one()
    return DistrictSummaryTotals(
        detected=row.detected,
        accepted=row.accepted,
        on_review=row.on_review,
        open=row.open,
        returned_now=row.returned_now,
        percent_text=percent_label(row.accepted, row.detected),
        returns_count=returns_count,
        avg_days_to_accept=_days(row.accept_seconds),
    )


async def _dynamics(
    db: AsyncSession,
    period: Period,
    conds: list,
    today: date,
) -> tuple[DynamicsUnit, list[DynamicsBucket]]:
    per_day = (
        await db.execute(
            select(
                _MSK_DAY.label("day"),
                func.count().label("detected"),
                func.count().filter(Card.status == "accepted").label("accepted"),
            )
            .where(*conds)
            .group_by(_MSK_DAY)
        )
    ).all()

    if period.date_from is not None and period.date_to is not None:
        start, end = period.date_from, period.date_to
    else:
        # «За всё время» — с первой карточки района; без карточек динамики нет.
        if not per_day:
            return "day", []
        start = min(r.day for r in per_day)
        end = max(today, max(r.day for r in per_day))

    unit, ranges = dynamics_buckets(start, end)
    counts = [[0, 0] for _ in ranges]
    for r in per_day:
        if unit == "day":
            idx = (r.day - start).days
        else:
            idx = (r.day.year - start.year) * 12 + r.day.month - start.month
        if 0 <= idx < len(counts):
            counts[idx][0] += r.detected
            counts[idx][1] += r.accepted
    return unit, [
        DynamicsBucket(
            label=label,
            date_from=b_from,
            date_to=b_to,
            detected=detected,
            accepted=accepted,
            percent_text=percent_label(accepted, detected),
        )
        for (b_from, b_to, label), (detected, accepted) in zip(ranges, counts)
    ]


async def _oldest_open(db: AsyncSession, district_id, today: date) -> list[OldestOpenCard]:
    cards = (
        (
            await db.execute(
                select(Card)
                .where(Card.district_id == district_id, Card.status.in_(OPEN_STATUSES))
                .order_by(Card.created_at.asc(), Card.number.asc())
                .limit(OLDEST_OPEN_LIMIT)
            )
        )
        .scalars()
        .all()
    )
    return [
        OldestOpenCard(
            id=c.id,
            label=card_label(c.number),
            address=c.address,
            status=cast(Literal["detected", "on_review", "accepted", "returned"], c.status),
            created_at=c.created_at,
            age_days=max((today - c.created_at.astimezone(MSK).date()).days, 0),
        )
        for c in cards
    ]


async def build_district_summary(
    db: AsyncSession,
    district: District,
    period: Period,
    now: datetime | None = None,
) -> DistrictSummaryOut:
    today = msk_today(now or datetime.now(timezone.utc))
    conds = _period_conds(district.id, period)
    unit, dynamics = await _dynamics(db, period, conds, today)
    return DistrictSummaryOut(
        district=DistrictOut(id=district.id, name=district.name),
        period=PeriodOut(
            kind=period.kind, date_from=period.date_from, date_to=period.date_to, label=period.label
        ),
        totals=await _totals(db, conds),
        dynamics_unit=unit,
        dynamics=dynamics,
        oldest_open=await _oldest_open(db, district.id, today),
    )


# ── Excel ───────────────────────────────────────────────────────

NCOLS = 5
KPI_HEADER = ["Показатель", "Значение"]
OLDEST_HEADER = ["Карточка", "Адрес", "Статус", "Выявлено", "Ждёт, дней"]


def district_xlsx_title(name: str) -> str:
    return f"Свод по люкам — {name}"


def _section(ws, text: str) -> None:
    ws.append([])
    safe_append(ws, [text])
    style_merged_label(ws, ws.max_row, NCOLS, header=True, fill=False)


def _table_header(ws, header: list[str]) -> None:
    safe_append(ws, header)
    style_header_row(ws, ws.max_row, len(header))


def _data_row(ws, row: list) -> None:
    safe_append(ws, row)
    style_data_row(ws, ws.max_row, len(row))


def _note(ws, text: str) -> None:
    safe_append(ws, [text])
    style_merged_label(ws, ws.max_row, NCOLS)
    ws.row_dimensions[ws.max_row].height = 32


def district_summary_xlsx(summary: DistrictSummaryOut) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Свод по району"
    t = summary.totals

    safe_append(ws, [district_xlsx_title(summary.district.name)])
    style_merged_label(ws, ws.max_row, NCOLS, header=True, fill=False)
    safe_append(ws, [f"Период: {summary.period.label}"])
    style_merged_label(ws, ws.max_row, NCOLS)

    _section(ws, "Показатели")
    _table_header(ws, KPI_HEADER)
    for label, value in (
        ("Выявлено (неудовлетворительные ОЛХ)", t.detected),
        ("Исправлено", t.accepted),
        ("На проверке", t.on_review),
        ("Не исправлено (выявлено и возвращено)", t.open),
        ("% исправления", t.percent_text),
        ("Возвратов на доработку", t.returns_count),
        (
            "Среднее время до приёмки, дней",
            t.avg_days_to_accept if t.avg_days_to_accept is not None else "—",
        ),
    ):
        _data_row(ws, [label, value])

    _section(ws, "Динамика")
    unit_header = "День" if summary.dynamics_unit == "day" else "Месяц"
    _table_header(ws, [unit_header, "Выявлено", "Исправлено", "% исправления"])
    for b in summary.dynamics:
        label = b.date_from.strftime("%d.%m.%Y") if summary.dynamics_unit == "day" else b.label
        _data_row(ws, [label, b.detected, b.accepted, b.percent_text])
    if not summary.dynamics:
        _note(ws, "Карточек пока нет")

    _section(ws, "Дольше всех ждут исправления")
    _table_header(ws, OLDEST_HEADER)
    for c in summary.oldest_open:
        _data_row(
            ws,
            [
                c.label,
                c.address,
                STATUS_TEXT[c.status],
                c.created_at.astimezone(MSK).strftime("%d.%m.%Y"),
                c.age_days,
            ],
        )
    if not summary.oldest_open:
        _note(ws, "Неисправленных карточек нет")

    ws.append([])
    _note(ws, FOOTNOTE)
    _note(ws, PERIOD_NOTE)

    for letter, width in zip("ABCDE", (38, 40, 16, 14, 14)):
        ws.column_dimensions[letter].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
